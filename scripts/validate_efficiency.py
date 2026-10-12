"""Check archived efficiency experiments, score constraints and publication."""
import argparse,ast,base64,hashlib,io,json,re,zipfile
from pathlib import Path
import numpy as np
ROOT=Path(__file__).resolve().parents[1]


def main(publication=False):
    report=json.loads((ROOT/'reports/efficiency_results.json').read_text(encoding='utf-8'))
    stages=['efficiency_scalar_v3','efficiency_geometry_v2','efficiency_vdp_final_0','efficiency_vdp_final_1']
    records=[]
    for stage in stages:
        folder=ROOT/f'results/kaggle_{stage}'
        assert json.loads((folder/'completion.json').read_text())['status']=='complete'
        manifest=json.loads((folder/'manifest.json').read_text())
        notebook=json.loads((ROOT/f'kaggle/{stage}/wsindy_{stage}.ipynb').read_text())
        payload=None
        for cell in notebook['cells']:
            for statement in ast.parse(''.join(cell['source'])).body:
                if isinstance(statement,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='payload' for t in statement.targets):
                    payload=ast.literal_eval(statement.value)
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(payload))) as bundle:
            assert bundle.testzip() is None
            fingerprint=hashlib.sha256()
            for directory in ('src','scripts','tests'):
                for name in sorted(n for n in bundle.namelist() if n.startswith(directory+'/') and n.endswith('.py')):
                    fingerprint.update(name.encode());fingerprint.update(bundle.read(name))
            assert fingerprint.hexdigest()==manifest['source_sha256']
            assert json.loads(bundle.read('configs/efficiency.json'))==manifest['protocol']
        path=ROOT/f'results/kaggle_{stage}_download/wsindy_{stage}_outputs.zip'
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            for name in archive.namelist():
                assert archive.read(name)==(folder/name).read_bytes()
        records.append(dict(stage=stage,output_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),execution_fingerprint_verified=True))
    assert report['scalar_fits']==33000 and report['vdp_fits']==1536
    rows=json.loads((ROOT/'results/kaggle_efficiency_scalar_v3/rows.json').read_text())
    assert len(rows)==33000
    for r in report['scalar']+report['vdp']:
        assert abs(r['relative_rmse']**2-r['relative_bias']**2-r['relative_dispersion']**2)<1e-12
        assert r['n']==(1000 if 'sigma' in r else 128)
    refs=report['scalar_reference']
    assert max(abs(r['initial_annihilation']) for r in refs)<1e-6
    assert max(abs(r['clean_quadrature_error']) for r in refs)<1e-6
    assert all(0<r['weak_efficiency']<=1+1e-5 for r in refs)
    for sigma in (.1,.02,.005):
        variances=[r['variance'] for r in refs if r['sigma']==sigma]
        assert np.all(np.diff(variances)<0)
    geometry=report['geometry']
    assert geometry['exact_covariance_fixture_relative_error']<1e-10
    bound=report['vdp_weak_relative_rmse_bound']
    for r in geometry['results']:
        C=np.array(r['actual_ou_covariance']);assert np.linalg.eigvalsh(C).min()>0
        assert r['relative_rmse']>=bound*(1-1e-5)
        assert r['normalization_error']<1e-8
        assert r['clean_quadrature_relative_error']<bound/1000
    matrices=[np.array(r['actual_ou_covariance']) for r in geometry['results'] if r['weight_model']=='ou']
    for left,right in zip(matrices,matrices[1:]):assert np.linalg.eigvalsh(left-right).min()>-1e-9
    assert all(r['converged']==128 for r in report['vdp'])
    result=dict(completed=True,stages=records,expected_seed_counts_and_finite_estimates=True,
        source_snapshots_and_protocols_verified=True,bias_variance_identity=True,
        scalar_annihilation_quadrature_and_class_bound=True,exact_covariance_fixture=True,
        matrix_covariance_positive_and_gls_normalized=True,publication_checked=publication)
    if publication:
        from docx import Document
        import pymupdf
        text=(ROOT/'manuscript/article_revised.md').read_text(encoding='utf-8')
        supp=(ROOT/'manuscript/supplement.md').read_text(encoding='utf-8')
        assert not re.search(r'[\u0400-\u04ff]',text+supp), 'Corrupted non-English text in manuscript'
        assert not re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',text+supp), 'Control character in manuscript'
        assert all(f'### 4.{i}.' in text for i in range(1,3))
        assert all(f'### 5.{i}.' in text for i in range(1,5))
        assert '**Lemma 1 ' in text
        assert all(f'**Proposition {i} ' in text for i in range(1,4))
        assert '**Theorem 2 ' in text and '**Corollary 1 ' in text
        assert 'sequential' in text and 'uniform integrability' in text
        assert 'malyshevdanil' not in supp and 'credentials' not in supp
        assert not re.search(r'abandoned grid diagnostic|initial custom-STLS pilot|historical|earlier 4,000-path experiment|fresh validation',text,re.I)
        assert len(Document(ROOT/'manuscript/article_revised.docx').tables)==6
        assert len(Document(ROOT/'manuscript/article_revised.docx').inline_shapes)==5
        assert len(Document(ROOT/'manuscript/supplement.docx').tables)==17
        assert '**Theorem 3 ' in text and 'Independent researcher, Almaty, Kazakhstan' in text
        tables=re.findall(r'^Table (S\d+)\.',supp,flags=re.M)
        assert tables==[f'S{i}' for i in range(1,18)]
        for name in ('article_revised','supplement','weak_efficiency_theory'):
            with pymupdf.open(ROOT/f'manuscript/{name}.pdf') as pdf:
                prose='\n'.join(p.get_text() for p in pdf)
                assert len(prose)>2000
                if name=='article_revised':
                    assert all(f'Table {i}.' in prose for i in range(1,7))
                    assert 'Data availability' in prose and 'Sharp information' in prose
                if name=='supplement':assert all(f'Table S{i}.' in prose for i in range(1,18))
        result['theorem_conditions_sequential_numbering_main_and_supplement_verified']=True
    (ROOT/'reports/efficiency_validation.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--publication',action='store_true');main(parser.parse_args().publication)
