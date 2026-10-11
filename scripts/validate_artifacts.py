"""Inspect built files and packed sources without executing empirical runs."""
import ast
import base64
import io
import json
import hashlib
import re
from pathlib import Path
import zipfile
import nbformat
from docx import Document
from lxml import etree
import pymupdf

ROOT = Path(__file__).resolve().parents[1]


def validate():
    records = []
    for path in sorted((ROOT/'kaggle').glob('*/*.ipynb')):
        notebook = nbformat.read(path, as_version=4)
        nbformat.validate(notebook)
        payload = None
        for cell in notebook.cells:
            if cell.cell_type == 'code':
                syntax = ast.parse(cell.source)
                for statement in syntax.body:
                    if isinstance(statement, ast.Assign) and any(
                            isinstance(t, ast.Name) and t.id == 'payload' for t in statement.targets):
                        payload = ast.literal_eval(statement.value)
                if payload is None:
                    for node in ast.walk(syntax):
                        if isinstance(node,ast.Call) and isinstance(node.func,ast.Attribute) and node.func.attr=='b64decode':
                            if node.args and isinstance(node.args[0],ast.Constant) and isinstance(node.args[0].value,str):
                                payload=node.args[0].value
        assert payload is not None
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(payload))) as archive:
            assert archive.testzip() is None
            differences = []
            for entry in archive.infolist():
                assert not Path(entry.filename).is_absolute() and '..' not in Path(entry.filename).parts
                assert Path(entry.filename).name not in {'access_token', 'kaggle.json', '.env'}
                expected = ROOT/entry.filename
                if not expected.exists() or archive.read(entry.filename) != expected.read_bytes():
                    differences.append(entry.filename)
                if entry.filename.endswith('.py'):
                    ast.parse(archive.read(entry.filename).decode('utf-8'))
            packed = len(archive.infolist())
            fingerprint=hashlib.sha256()
            for folder in ('src','scripts','tests'):
                for name in sorted(n for n in archive.namelist() if n.startswith(folder+'/') and n.endswith('.py')):
                    fingerprint.update(name.encode())
                    fingerprint.update(archive.read(name))
            stage=path.parent.name
            manifest_path=ROOT/f'results/kaggle_{stage}/manifest.json'
            if stage=='smoke':
                manifest_path=ROOT/'results/kaggle_smoke_download/results_smoke/manifest.json'
            snapshot_verified=False
            if manifest_path.exists():
                manifest=json.loads(manifest_path.read_text())
                if stage.startswith('adaptive_'):
                    names=['scripts/run_full_adaptive.py','configs/full_adaptive.json','configs/author_pins.json']
                    names += [n for n in archive.namelist() if n.startswith('src/') and n.endswith('.py')]
                    names += [n for n in archive.namelist() if n.startswith('src/author_compat/') and n.endswith('.m')]
                    calculated=hashlib.sha256(b''.join(n.encode()+archive.read(n) for n in sorted(names))).hexdigest()
                    assert calculated==manifest['source_fingerprint'],f'Execution source hash mismatch: {stage}'
                else:
                    calculated=(hashlib.sha256(archive.read('scripts/run_extended_validation.py')).hexdigest()
                                if stage=='extended' else fingerprint.hexdigest())
                    assert calculated==manifest['source_sha256'], f'Execution source hash mismatch: {stage}'
                snapshot_verified=True
        metadata = json.loads((path.parent/'kernel-metadata.json').read_text())
        assert metadata['is_private'] is True
        assert metadata['enable_gpu'] is False
        assert metadata['enable_internet'] is path.parent.name.startswith('adaptive_')
        # Completed notebooks are immutable execution snapshots, not regenerated
        # after local manuscript/postprocessing changes. Record differences openly.
        records.append({'file': str(path.relative_to(ROOT)), 'notebook_valid': True,
                        'packed_files': packed, 'workspace_differences': differences,
                        'execution_snapshot_preserved': True,
                        'downloaded_manifest_matches_snapshot':snapshot_verified})
    for path in (ROOT/'manuscript/article_en.docx', ROOT/'manuscript/article_revised.docx',
                 ROOT/'manuscript/supplement.docx',ROOT/'manuscript/theory.docx',
                 ROOT/'reports/research_progress_RU.docx',ROOT/'reports/q1_revision_results_RU.docx'):
        doc = Document(path)
        minimum_paragraphs=20 if path.parent.name=='manuscript' else 10
        assert len(doc.paragraphs)>minimum_paragraphs, f'Truncated document: {path.name}'
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            xml=etree.fromstring(archive.read('word/document.xml'))
            equations=int(xml.xpath('count(//m:oMath)',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}))
            if path.parent.name=='manuscript':
                assert equations>=5
                prose=' '.join(xml.xpath('//w:t/text()',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))
                assert not re.search(r'\\(?:frac|int|sum|widehat|Sigma|phi|mathcal|mathrm)\b',prose)
        records.append({'file': str(path.relative_to(ROOT)), 'docx_reopened': True,
                        'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables),
                        'figures': len(doc.inline_shapes),'native_omml_equations':equations})
    for name in ('article_revised','supplement','information_theory'):
        path=ROOT/f'manuscript/{name}.pdf'
        with pymupdf.open(path) as document:
            figure_count=0
            for page in document:
                for block in page.get_text('dict')['blocks']:
                    if block['type']==1:
                        rectangle=pymupdf.Rect(block['bbox'])
                        assert page.rect.contains(rectangle), f'Image outside PDF page: {name}'
                        figure_count+=1
            if name=='article_revised':
                source=(ROOT/'manuscript/article_revised.md').read_text(encoding='utf-8')
                assert figure_count==len(re.findall(r'^!\[Figure \d+\.',source,flags=re.M))
                text='\n'.join(page.get_text() for page in document)
                manuscript=(ROOT/'manuscript/article_revised.md').read_text(encoding='utf-8')
                tables=re.findall(r'^Table (\d+)\.',manuscript,flags=re.M)
                assert tables==[str(i) for i in range(1,len(tables)+1)]
                assert all(f'Table {i}.' in text for i in tables)
            records.append({'file':str(path.relative_to(ROOT)),'pdf_reopened':True,
                            'pages':len(document),'figures_within_page':figure_count})
    target = ROOT/'reports/artifact_validation.json'
    target.write_text(json.dumps(records, indent=2), encoding='utf-8')
    print(json.dumps(records, ensure_ascii=False))


if __name__ == '__main__':
    validate()
