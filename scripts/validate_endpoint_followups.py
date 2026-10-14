"""Validate compact endpoint archives, frozen sources and current publication artifacts."""
import ast,base64,hashlib,io,json,re,sys,zipfile
from pathlib import Path
import nbformat
import pymupdf
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))


def main():
    from scripts.collect_stopped_runs import snapshot
    dest=ROOT/'data/endpoint_followups';manifest=json.loads((dest/'manifest.json').read_text())
    for item in manifest['files']:
        data=(dest/item['path']).read_bytes();assert len(data)==item['bytes']
        assert hashlib.sha256(data).hexdigest()==item['sha256']
    audits=[]
    for stage in manifest['stages']:
        path=ROOT/f'kaggle/{stage}/wsindy_{stage}.ipynb';nb=nbformat.read(path,as_version=4);nbformat.validate(nb)
        meta=json.loads((path.parent/'kernel-metadata.json').read_text())
        assert meta['is_private'] and not meta['enable_gpu']
        assert meta['enable_internet']==stage.startswith('support_wald_')
        payload=None
        for cell in nb.cells:
            if cell.cell_type!='code':continue
            for node in ast.parse(cell.source).body:
                if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='payload' for t in node.targets):payload=ast.literal_eval(node.value)
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(payload))) as z:
            assert z.testzip() is None
            for name in z.namelist():
                assert not Path(name).is_absolute() and '..' not in Path(name).parts
                assert Path(name).name not in ['access_token','kaggle.json','.env']
                assert not re.search(rb'KGAT_[A-Za-z0-9]{20,}',z.read(name)),f'Token-shaped bytes in {name}'
                if name.endswith('.py'):ast.parse(z.read(name).decode('utf-8'))
        m=json.loads((dest/stage/'manifest.json').read_text());audit=snapshot(stage,m)
        assert json.loads((dest/stage/'completion.json').read_text())['status']=='complete'
        audits.append({'stage':stage,**audit})
    docs=[]
    for stem in ['article_revised','supplement']:
        document=Document(ROOT/f'manuscript/{stem}.docx')
        assert len(document.paragraphs)>100
        docs.append({'name':stem,'tables':len(document.tables),'figures':len(document.inline_shapes)})
    pdfs=[]
    for stem in ['article_ejs','article_revised','supplement']:
        with pymupdf.open(ROOT/f'manuscript/{stem}.pdf') as document:
            content='\n'.join(p.get_text() for p in document)
            assert 'Table 7.' in content if stem!='supplement' else 'Table S1.' in content
            for page in document:
                for block in page.get_text('dict')['blocks']:
                    if block['type']==1:assert page.rect.contains(pymupdf.Rect(block['bbox']))
            pdfs.append({'name':stem,'pages':len(document)})
    ejs=(ROOT/'manuscript/article_ejs.tex').read_text(encoding='utf-8')
    assert r'\documentclass[ejs]{imsart}' in ejs
    assert '62F12' in ejs and '62M10' in ejs and '34A55' in ejs
    main_text=(ROOT/'manuscript/article_revised.md').read_text(encoding='utf-8')
    assert '**Theorem 2 (known-initial extended influence class)' in main_text
    assert '**Proposition 3 (guarded fixed-family weak GLS)' in main_text
    assert '**Theorem 3 (constructive attainability under a jet condition)' in main_text
    assert '### 7.5. Prespecified testing versus sparse inclusion' not in main_text
    assert 'S7.2' in main_text
    supplement_text=(ROOT/'manuscript/supplement.md').read_text(encoding='utf-8')
    assert re.findall(r'^## S(\d+)\.',supplement_text,re.M)==[str(i) for i in range(1,10)]
    assert re.findall(r'^Table S(\d+)\.',supplement_text,re.M)==[str(i) for i in range(1,9)]
    assert re.findall(r'^!\[Figure S(\d+)\.',supplement_text,re.M)==['1','2','3']
    assert 'targeted selection' not in main_text
    assert 'remain supplementary extensions' not in main_text
    assert 'before any new endpoint outcomes' in main_text
    differences=json.loads((ROOT/'reports/support_wald_paired_intervals.json').read_text())
    assert differences['all_1600_paired_outcomes_retained'] and differences['counts_unchanged']
    assert len(differences['groups'])==16 and all(g['ci95'][0]<g['ci95'][1] for g in differences['groups'])
    for p in [ROOT/'manuscript/article_revised.md',ROOT/'manuscript/supplement.md']:
        text=p.read_text(encoding='utf-8');assert not re.search('KGAT_|credentials|access_token',text)
    result={'archive_files':len(manifest['files']),'sha256_all_verified':True,'frozen_source_audits':audits,
            'documents':docs,'pdfs':pdfs,'ejs_metadata_verified':True,'public_identifier_for_current_revision_pending':True}
    (ROOT/'reports/endpoint_followup_validation.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result))


if __name__=='__main__':main()
